import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from openpyxl import load_workbook
from datetime import datetime
from oop import Validate
import oop
print(oop.__file__)
# ===============================
# Load Excel Workbook
# ===============================
workbook = load_workbook("lab_tests.xlsx")

test_sheet = workbook["tests"]
doctor_sheet = workbook["dr_details"]


# ===============================
# Read Tests
# ===============================
tests = []

for row in test_sheet.iter_rows(min_row=2, values_only=True):
    if row[0] and row[1]:
        tests.append(f"{row[0]} | {row[1]}")

# ===============================
# Read Doctors
# ===============================
doctors = []

for row in doctor_sheet.iter_rows(min_row=2, values_only=True):
    if row[0]:
        doctors.append(str(row[0]))

# ===============================
# Main Window
# ===============================
#from tkinter import messagebox

def add_record():
    test_code = test_combo.get().split(" | ")[0]
    
    if test_combo.get() == "":
        messagebox.showerror("Error", "Please select a test.")
        return

    if title_combo.get() == "":
        messagebox.showerror("Error", "Please select a title.")
        return

    if name_entry.get() == "":
        messagebox.showerror("Error", "Please enter patient name.")
        return

    if nic_entry.get() == "":
        messagebox.showerror("Error", "Please enter NIC / Passport.")
        return

    if patient_combo.get() == "":
        messagebox.showerror("Error", "Please select patient type.")
        return
    if doctor_combo.get() == "":
        messagebox.showerror("Error", "Please select doctor.")
        return

    # -----------------------------
    # Date Validation
    # -----------------------------
    if not Validate.validate_date(date_entry.get()):
        messagebox.showerror(
            "Invalid Date",
            "Date must be in DD-MM-YYYY format."
        )
        return

    # -----------------------------
    # NIC Validation
    # -----------------------------
    if not Validate.validate_nic(nic_entry.get()):
        messagebox.showerror(
            "Invalid NIC / Passport",
            "Enter a valid NIC or Passport number."
        )
        return

    # -----------------------------
    # Patient Type Validation
    # -----------------------------
    test_code = test_combo.get().split(" | ")[0]

    if not Validate.validate_patient_type(
            test_code,
            patient_combo.get()):

        messagebox.showerror(
            "Invalid Patient Type",
            "Selected test is not available for this patient type."
        )
        return

    test_code, test_name = test_combo.get().split(" | ")

    patient_name = title_combo.get() + " " + name_entry.get()

    tree.insert(
        "",
        "end",
        values=(
            test_code,
            test_name,
            patient_name,
            nic_entry.get(),
            patient_combo.get(),
            doctor_combo.get()
        )
    )

    # Clear form after adding
    test_combo.set("")
    title_combo.set("")
    name_entry.delete(0, tk.END)
    nic_entry.delete(0, tk.END)
    patient_combo.set("")
    doctor_combo.set("")
def remove_record():

    selected = tree.selection()

    if not selected:
        messagebox.showwarning("Warning", "Please select a record to remove.")
        return

    confirm = messagebox.askyesno(
        "Confirm",
        "Are you sure you want to remove the selected record?"
    )

    if confirm:
        tree.delete(selected)
def save_records():

    if len(tree.get_children()) == 0:
        messagebox.showinfo("No Records", "There are no records to save.")
        return

    confirm = messagebox.askyesno(
        "Confirm Save",
        "Do you want to save all records?"
    )

    if not confirm:
        return

    workbook = load_workbook("lab_tests.xlsx")

    for item in tree.get_children():

        values = tree.item(item)["values"]

        test_code = values[0]

        if test_code in workbook.sheetnames:
            sheet = workbook[test_code]
        else:
            sheet = workbook.create_sheet(test_code)

            sheet.append([
                "Date",
                "Patient Name",
                "Patient Type",
                "Patient NIC/PP",
                "Test Name",
                "Dr / Consultant"
            ])

        sheet.append([
            date_entry.get(),
            values[2],
            values[4],
            values[3],
            values[1],
            values[5]
        ])

    workbook.save("lab_tests.xlsx")
    workbook.close()
    messagebox.showinfo(
        "Success",
        "Records saved successfully!"
    )
def clear_records():

    if len(tree.get_children()) == 0:
        messagebox.showinfo(
            "No Records",
            "There are no records to clear."
        )
        return

    confirm = messagebox.askyesno(
        "Confirm",
        "Do you want to clear all records?"
    )

    if confirm:
        for item in tree.get_children():
            tree.delete(item)
root = tk.Tk()
root.title("Hospital Lab Test Log")
root.geometry("1000x700")

# ===============================
# Frame 1
# ===============================
frame1 = tk.LabelFrame(root, text="Patient Details", padx=10, pady=10)
frame1.pack(fill="x", padx=10, pady=10)

# Date
tk.Label(frame1, text="Date").grid(row=0, column=0, padx=5, pady=5)

date_var = tk.StringVar(value=datetime.today().strftime("%d-%m-%Y"))

date_entry = tk.Entry(frame1, textvariable=date_var, width=20)
date_entry.grid(row=0, column=1, padx=5, pady=5)

# Test
tk.Label(frame1, text="Test").grid(row=1, column=0, padx=5, pady=5)

test_combo = ttk.Combobox(frame1, values=tests, width=45)
test_combo.grid(row=1, column=1, padx=5, pady=5)

# Title
tk.Label(frame1, text="Title").grid(row=2, column=0, padx=5, pady=5)

title_combo = ttk.Combobox(
    frame1,
    values=["Mr", "Mrs", "Miss", "Dr"],
    width=10
)
title_combo.grid(row=2, column=1, sticky="w", padx=5, pady=5)

# Patient Name
tk.Label(frame1, text="Patient Name").grid(row=3, column=0, padx=5, pady=5)

name_entry = tk.Entry(frame1, width=40)
name_entry.grid(row=3, column=1, padx=5, pady=5)

# NIC
tk.Label(frame1, text="NIC / Passport").grid(row=4, column=0, padx=5, pady=5)

nic_entry = tk.Entry(frame1, width=40)
nic_entry.grid(row=4, column=1, padx=5, pady=5)

# Patient Type
tk.Label(frame1, text="Patient Type").grid(row=5, column=0, padx=5, pady=5)

patient_combo = ttk.Combobox(
    frame1,
    values=["Inpatient", "Outpatient"],
    width=20
)
patient_combo.grid(row=5, column=1, sticky="w", padx=5, pady=5)

# Doctor
tk.Label(frame1, text="Doctor").grid(row=6, column=0, padx=5, pady=5)

doctor_combo = ttk.Combobox(frame1, values=doctors, width=40)
doctor_combo.grid(row=6, column=1, padx=5, pady=5)


# ===============================
# Frame 2
# ===============================

frame2 = tk.LabelFrame(root, text="Lab Test Records", padx=10, pady=10)
frame2.pack(fill="both", expand=True, padx=10, pady=10)

columns = (
    "Test Code",
    "Test Name",
    "Patient Name",
    "Patient NIC/PP",
    "Patient Type",
    "Dr / Consultant"
)

tree = ttk.Treeview(frame2, columns=columns, show="headings")

for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=150, anchor="center")

tree.pack(fill="both", expand=True)


# ===============================
# Buttons
# ===============================

add_btn = tk.Button(
    frame1,
    text="Add",
    width=12,
    command=add_record
)
add_btn.grid(row=7, column=0, padx=5, pady=10)

remove_btn = tk.Button(
    frame1,
    text="Remove",
    width=12,
    command=remove_record
)
remove_btn.grid(row=7, column=1, sticky="w", padx=5, pady=10)


# ===============================
# Frame 3
# ===============================

frame3 = tk.Frame(root)


# ===============================
# Frame 3
# ===============================

frame3 = tk.Frame(root)
frame3.pack(fill="x", padx=10, pady=10)

save_btn = tk.Button(
    frame3,
    text="Save",
    width=15,
    command=save_records
)
save_btn.pack(side="left", padx=10)

clear_btn = tk.Button(
    frame3,
    text="Clear",
    width=15,
    command=clear_records
)
clear_btn.pack(side="left", padx=10)

# ===============================
# Run Program
# ===============================

root.mainloop()