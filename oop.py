import re
from datetime import datetime
from openpyxl import load_workbook


class Validate:

    # Validate Date
    @staticmethod
    def validate_date(date):
        try:
            datetime.strptime(date, "%d-%m-%Y")
            return True
        except ValueError:
            return False

    # Validate NIC / Passport
    @staticmethod
    def validate_nic(nic):
        pattern1 = r'^\d{12}$'
        pattern2 = r'^\d{9}[VvXx]$'
        pattern3 = r'^[NP]\d{7}$'

        if re.match(pattern1, nic):
            return True

        if re.match(pattern2, nic):
            return True

        if re.match(pattern3, nic):
            return True

        return False

    # Validate Patient Type
    @staticmethod
    def validate_patient_type(test_code, patient_type):

        workbook = load_workbook("lab_tests.xlsx")
        sheet = workbook["tests"]

        try:
            for row in sheet.iter_rows(min_row=2, values_only=True):

                if row[0] is None:
                    continue

                if str(row[0]).strip() == str(test_code).strip():

                    allowed = str(row[2]).strip().lower()

                    if allowed == "in/out":
                        return True

                    elif allowed == "in":
                        return patient_type.lower() == "inpatient"

                    elif allowed == "out":
                        return patient_type.lower() == "outpatient"

                    else:
                        return False

            return False

        finally:
            workbook.close()